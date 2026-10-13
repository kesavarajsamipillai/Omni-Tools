"""
Document Converter Module (Completely Offline)
Features:
- PDF to DOCX (via pdf2docx)
- DOCX to PDF (via Microsoft Word automation with pure-python ReportLab fallback)
- Text to PDF (via ReportLab)
"""

import os
from typing import Optional


def pdf_to_docx(input_pdf: str, output_docx: str, start_page: int = 0, end_page: Optional[int] = None) -> dict:
    """Converts a PDF file to editable Microsoft Word DOCX format offline."""
    if not os.path.exists(input_pdf):
        return {"success": False, "error": f"File not found: {input_pdf}"}

    try:
        from pdf2docx import Converter

        os.makedirs(os.path.dirname(os.path.abspath(output_docx)), exist_ok=True)
        cv = Converter(input_pdf)
        cv.convert(output_docx, start=start_page, end=end_page)
        cv.close()

        return {"success": True, "output": output_docx}
    except Exception as e:
        return {"success": False, "error": str(e)}


def docx_to_pdf(input_docx: str, output_pdf: str) -> dict:
    """
    Converts DOCX to PDF.
    First attempts native MS Word automation (docx2pdf) for perfect layout rendering.
    If MS Word is not installed, seamlessly falls back to pure-Python DOCX parsing with ReportLab!
    """
    if not os.path.exists(input_docx):
        return {"success": False, "error": f"File not found: {input_docx}"}

    os.makedirs(os.path.dirname(os.path.abspath(output_pdf)), exist_ok=True)

    # Method 1: Try docx2pdf (Requires Word installed on Windows)
    try:
        from docx2pdf import convert
        convert(input_docx, output_pdf)
        if os.path.exists(output_pdf) and os.path.getsize(output_pdf) > 0:
            return {"success": True, "output": output_pdf, "engine": "MS Word Native"}
    except Exception:
        pass  # Fall back to pure Python

    # Method 2: Pure Python fallback using python-docx + ReportLab
    try:
        import docx
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib import colors

        doc = docx.Document(input_docx)
        pdf_doc = SimpleDocTemplate(
            output_pdf,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()
        normal_style = styles["Normal"]
        heading_style = styles["Heading1"]
        h2_style = styles["Heading2"]
        story = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                story.append(Spacer(1, 6))
                continue
            
            # Escape XML entities for ReportLab Paragraph
            safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

            if p.style.name.startswith("Heading 1"):
                story.append(Paragraph(safe_text, heading_style))
                story.append(Spacer(1, 10))
            elif p.style.name.startswith("Heading 2"):
                story.append(Paragraph(safe_text, h2_style))
                story.append(Spacer(1, 8))
            else:
                story.append(Paragraph(safe_text, normal_style))
                story.append(Spacer(1, 6))

        for tbl in doc.tables:
            table_data = []
            for row in tbl.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            if table_data:
                t = Table(table_data)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
                    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                    ('FONTSIZE', (0,0), (-1,-1), 9),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.grey)
                ]))
                story.append(t)
                story.append(Spacer(1, 12))

        if not story:
            story.append(Paragraph("(Empty Document)", normal_style))

        pdf_doc.build(story)
        return {"success": True, "output": output_pdf, "engine": "Pure Python ReportLab"}
    except Exception as e:
        return {"success": False, "error": f"DOCX to PDF conversion failed: {str(e)}"}


def text_to_pdf(input_txt: str, output_pdf: str, font_size: int = 10) -> dict:
    """Converts a plain text file into a clean formatted PDF."""
    if not os.path.exists(input_txt):
        return {"success": False, "error": f"File not found: {input_txt}"}

    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

        with open(input_txt, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()

        os.makedirs(os.path.dirname(os.path.abspath(output_pdf)), exist_ok=True)
        pdf = SimpleDocTemplate(
            output_pdf,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()
        mono_style = ParagraphStyle(
            name="MonoText",
            fontName="Courier",
            fontSize=font_size,
            leading=font_size + 3
        )

        story = []
        for line in lines:
            safe_line = line.rstrip("\r\n").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace(" ", "&nbsp;")
            if not safe_line:
                story.append(Spacer(1, font_size))
            else:
                story.append(Paragraph(safe_line, mono_style))

        pdf.build(story)
        return {"success": True, "output": output_pdf, "lines_processed": len(lines)}
    except Exception as e:
        return {"success": False, "error": str(e)}
